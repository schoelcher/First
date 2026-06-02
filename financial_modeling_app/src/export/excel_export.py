"""Excel workbook export for the financial model."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def _write_kv(ws, rows: list[tuple[str, Any]], start_row: int = 1) -> None:
    for idx, (key, value) in enumerate(rows, start=start_row):
        ws.cell(idx, 1, key).font = Font(bold=True)
        ws.cell(idx, 2, value)


def _write_df(ws, df: pd.DataFrame, start_row: int = 1, start_col: int = 1) -> None:
    ws.cell(start_row, start_col, df.index.name or "index").font = Font(bold=True)
    for col_idx, col in enumerate(df.columns, start=start_col + 1):
        cell = ws.cell(start_row, col_idx, str(col))
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row_idx, (index, row) in enumerate(df.iterrows(), start=start_row + 1):
        ws.cell(row_idx, start_col, index)
        for col_idx, value in enumerate(row, start=start_col + 1):
            cell = ws.cell(row_idx, col_idx, None if pd.isna(value) else value)
            column_name = str(df.columns[col_idx - start_col - 1]).lower()
            if any(token in column_name for token in ["margin", "growth", "rate", "percent", "%", "upside"]):
                cell.number_format = "0.0%"
            elif "multiple" in column_name:
                cell.number_format = "0.0x"
            elif isinstance(value, (int, float)):
                cell.number_format = '#,##0.0'


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width, 10), 35)


def export_model_to_excel(
    ticker: str,
    company_name: str,
    historical: pd.DataFrame,
    assumptions: Any,
    forecast: pd.DataFrame,
    wacc_result: Any,
    dcf_result: Any,
    sensitivity_tables: dict[str, pd.DataFrame],
    warnings: list[str],
) -> bytes:
    """Build an Excel workbook and return it as bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    dcf_dict = dcf_result.to_dict() if hasattr(dcf_result, "to_dict") else dict(dcf_result)
    _write_kv(
        ws,
        [
            ("Ticker", ticker),
            ("Company", company_name),
            ("Generated UTC", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
            ("Enterprise Value", dcf_dict.get("enterprise_value")),
            ("Equity Value", dcf_dict.get("equity_value")),
            ("Implied Share Price", dcf_dict.get("implied_share_price")),
            ("Current Share Price", dcf_dict.get("current_share_price")),
            ("Upside / Downside", dcf_dict.get("upside_downside")),
            ("WACC", dcf_dict.get("wacc")),
            ("Terminal Method", dcf_dict.get("terminal_method")),
        ],
    )
    ws.cell(13, 1, "Warnings / Data Quality Notes").font = Font(bold=True)
    for idx, warning in enumerate(warnings, start=14):
        ws.cell(idx, 1, warning)

    for title, df in [("Historical Financials", historical), ("Forecast", forecast)]:
        sheet = wb.create_sheet(title)
        _write_df(sheet, df)

    assump_sheet = wb.create_sheet("Assumptions")
    assump_dict = assumptions.model_dump() if hasattr(assumptions, "model_dump") else assumptions.__dict__
    _write_kv(assump_sheet, list(assump_dict.items()))

    dcf_sheet = wb.create_sheet("DCF")
    wacc_dict = wacc_result.__dict__ if hasattr(wacc_result, "__dict__") else dict(wacc_result)
    _write_kv(dcf_sheet, list(wacc_dict.items()) + list(dcf_dict.items()))

    sens_sheet = wb.create_sheet("Sensitivity")
    row = 1
    for name, table in sensitivity_tables.items():
        sens_sheet.cell(row, 1, name).font = Font(bold=True, size=12)
        _write_df(sens_sheet, table, start_row=row + 1)
        row += len(table) + 4

    for sheet in wb.worksheets:
        _autosize(sheet)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
